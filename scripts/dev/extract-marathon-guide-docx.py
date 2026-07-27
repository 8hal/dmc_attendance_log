#!/usr/bin/env python3
"""Extract marathon guide DOCX → JSON (stdout). Not a runtime dependency."""
from __future__ import annotations

import json
import os
import sys

DEFAULT_DOCX = (
    "/tmp/marathon-guide-handoff/마라톤_100일_가이드_handoff/final/"
    "2026_마라톤_100일_수준별_훈련_가이드.docx"
)


def main() -> int:
    path = os.environ.get("GUIDE_DOCX", DEFAULT_DOCX)
    if not os.path.isfile(path):
        print(f"DOCX not found: {path}", file=sys.stderr)
        return 1

    try:
        from docx import Document
    except ImportError:
        print("python-docx is required: pip install python-docx", file=sys.stderr)
        return 1

    doc = Document(path)

    paragraphs = []
    headings = []
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = p.style.name if p.style is not None else ""
        entry = {"style": style, "text": p.text or ""}
        paragraphs.append(entry)
        if style.startswith("Heading"):
            level = 1
            parts = style.split()
            if len(parts) >= 2 and parts[-1].isdigit():
                level = int(parts[-1])
            headings.append({"level": level, "style": style, "text": p.text or ""})

    tables = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cells.append(cell.text or "")
            rows.append(cells)
        tables.append({"i": i, "rows": rows})

    payload = {
        "source": path,
        "headings": headings,
        "paragraphs": paragraphs,
        "paras": paragraphs,  # alias for earlier dump consumers
        "tables": tables,
    }
    json.dump(payload, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except BrokenPipeError:
        # Allow `... | head` without a traceback.
        try:
            sys.stdout.close()
        except Exception:
            pass
        raise SystemExit(0)
